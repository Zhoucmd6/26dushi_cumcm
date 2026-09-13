"""只读检查导出的模板：逐值、逐日期及事件合计，不用openpyxl写文件。"""
from pathlib import Path
import sys,json
from datetime import datetime
import numpy as np
from openpyxl import load_workbook
run=Path(sys.argv[1]);p=json.loads((run/'q3_workbook_payload.json').read_text(encoding='utf-8'))
wb=load_workbook(run/'deliverables/result3.xlsx',data_only=True)
assert wb.sheetnames==['计划购电量','调整购电量','充放电量','紧急购电量']
checks={}
for sheet,key in [('计划购电量','plan'),('调整购电量','adjusted')]:
    rows=list(wb[sheet].iter_rows(min_row=2,max_row=335,max_col=147,values_only=True))
    for i,row in enumerate(rows):
        assert row[0].date().isoformat()==p[key][i][0]
        np.testing.assert_allclose(row[1:],p[key][i][1:],rtol=0,atol=1e-6)
    assert wb[sheet]['B1'].value=='00:00-00:10';assert wb[sheet]['EO1'].value=='23:50-24:00'
    checks[sheet]={'dates':334,'ten_minute_values':334*144,'date_and_value_checks':'passed'}
for sheet,key,cols in [('充放电量','battery',6),('紧急购电量','emergency',3)]:
    rows=list(wb[sheet].iter_rows(min_row=2,max_row=len(p[key])+1,max_col=cols,values_only=True))
    for actual,expected in zip(rows,p[key]):
        for j,(a,e) in enumerate(zip(actual,expected)):
            if j==0 and e is not None:assert a.date().isoformat()==e
            elif isinstance(e,(int,float)):assert a is not None and abs(a-e)<1e-6
            else:assert a==e,(sheet,j,a,e)
    checks[sheet]={'rows':len(rows),'values':'passed'}
total=sum(wb['调整购电量'].cell(i,147).value for i in range(2,336))
assert abs(total-p['total_fee_yuan'])<1e-5
for s in wb:
    for row in s:
        assert all(c.data_type!='e' for c in row)
checks['total_fee_yuan']=total
(run/'workbook_qa/openpyxl_readonly_check.json').write_text(json.dumps(checks,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(checks,ensure_ascii=False))
