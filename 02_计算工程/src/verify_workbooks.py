"""Read-only saved-XLSX reconciliation, independent of Artifact Tool exporter."""
import argparse,json,hashlib
from datetime import datetime
from openpyxl import load_workbook
from run_q4 import ROOT,write_json

def main():
    p=argparse.ArgumentParser();p.add_argument('--run',default='full_q4_20260913');a=p.parse_args()
    run=ROOT/'outputs'/a.run;results=[]
    for mode,file in [('q42','result4-2.xlsx'),('q43','result4-3.xlsx')]:
        payload=json.loads((run/(mode+'_workbook_payload.json')).read_text('utf-8'))
        path=run/'deliverables'/file;wb=load_workbook(path,data_only=True)
        formulas=load_workbook(path,data_only=False,read_only=True);maximum=0.;numbers=0
        pairs=[('计划购电量','plan')]+([('调整购电量','adjusted')] if mode=='q43' else [])+[('充放电量','battery'),('紧急购电量','emergency')]
        assert wb.sheetnames==[name for name,_ in pairs]
        for name,key in pairs:
            s=wb[name];rows=payload[key]
            for i,row in enumerate(rows,2):
                for j,value in enumerate(row,1):
                    actual=s.cell(i,j).value
                    if j==1 and value is not None:
                        assert isinstance(actual,datetime) and actual.strftime('%Y-%m-%d')==value,(file,name,i,j,actual,value)
                    elif isinstance(value,(int,float)):
                        assert isinstance(actual,(int,float)),(file,name,i,j,actual)
                        error=abs(actual-value);maximum=max(maximum,error);numbers+=1
                        assert error<2e-5,(file,name,i,j,actual,value)
                    else:assert actual==value,(file,name,i,j,actual,value)
            if key in ('plan','adjusted'):
                assert [s.cell(1,j+2).value for j in range(144)]==payload['headers'][1:145]
                for i in (2,168,335):assert formulas[name].cell(i,146).value==f'=SUM(B{i}:EO{i})'
            for row in s.iter_rows():
                for cell in row:assert cell.data_type!='e',(file,name,cell.coordinate,cell.value)
        cash_sheet='计划购电量' if mode=='q42' else '调整购电量'
        cash=sum(wb[cash_sheet].cell(i,147).value for i in range(2,336))
        expected=json.loads((run/payload['experiment']/'summary.json').read_text('utf-8'))['cash_cost_yuan']
        assert abs(cash-expected)<2e-5
        results.append({'file':file,'sheets':wb.sheetnames,'numeric_cells_checked':numbers,'maximum_cell_error':maximum,
            'cash_cost_yuan':cash,'cash_difference_yuan':cash-expected,'sha256':hashlib.sha256(path.read_bytes()).hexdigest()})
        print(file,numbers,maximum,cash,flush=True)
    write_json(run/'workbook_verification.json',results)

if __name__=='__main__':main()
