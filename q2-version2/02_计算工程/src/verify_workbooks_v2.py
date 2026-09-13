"""只读核对实际导出的每个日期和数值；不会写入Excel。"""
from pathlib import Path
import sys,json
import numpy as np
from openpyxl import load_workbook


def verify(run):
    p = json.loads((run/'workbook_payload.json').read_text(encoding='utf-8'))
    checks = {}
    for number, names in [(1,['计划购电量','充放电量']), (2,['计划购电量','充放电量','紧急购电量'])]:
        w = load_workbook(run/f'deliverables/result{number}.xlsx',data_only=True)
        if w.sheetnames != names:
            raise ValueError('Template sheet topology changed')
        if number == 1:
            blocks = [('计划购电量',p['q1']['purchases']),('充放电量',p['q1']['battery'])]
        else:
            blocks = [('计划购电量',p['q2']['purchases']),('充放电量',p['q2']['battery']),('紧急购电量',p['q2']['emergency'])]
        for sheet, expected in blocks:
            actual = w[sheet].iter_rows(min_row=2,max_row=len(expected)+1,max_col=len(expected[0]),values_only=True)
            for index,(row,target) in enumerate(zip(actual,expected)):
                for col,(a,e) in enumerate(zip(row,target)):
                    if number == 2 and col == 0 and e is not None:
                        assert a.date().isoformat() == e, (sheet,index,col)
                    elif isinstance(e,(int,float)):
                        assert a is not None and abs(a-e) < 1e-6, (sheet,index,col,a,e)
                    else:
                        assert a == e, (sheet,index,col,a,e)
            checks[f'result{number}/{sheet}'] = {'rows':len(expected),'all_values_verified':True}
        for sheet in w:
            for row in sheet:
                assert all(c.data_type != 'e' for c in row)
        if number == 2:
            assert w['计划购电量']['B1'].value == '00:00-00:10'
            assert w['计划购电量']['EO1'].value == '23:50-24:00'
            assert 'version2' in w['计划购电量']['A338'].value
            checks['total_cash_cost_yuan'] = sum(w['计划购电量'].cell(i,147).value for i in range(2,336))
        w.close()
    (run/'workbook_qa/readonly_export_check.json').write_text(json.dumps(checks,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(checks,ensure_ascii=False,indent=2))


if __name__ == '__main__':
    verify(Path(sys.argv[1]).resolve())
