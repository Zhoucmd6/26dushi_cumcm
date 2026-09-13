"""独立审核完整轨迹与导出文件；不调用规划或运行回放函数。"""
from pathlib import Path
from datetime import date, datetime
import argparse
import csv
import hashlib
import json

import numpy as np
import openpyxl

from forecasting import read_dataset, interval_label, clock_text

ROOT=Path(__file__).resolve().parents[1]


def close(name,value,tolerance=1e-5):
    error=float(np.max(np.abs(value)))
    if not np.isfinite(error) or error>tolerance:
        raise AssertionError(f"{name}: {error}")
    return error


def audit_arrays(run,dataset,config):
    summaries=json.loads((run/"experiment_summaries.json").read_text(encoding="utf-8"))
    results={}
    for summary in summaries:
        name=summary["name"]
        folder=run/name
        with np.load(folder/"trajectories.npz") as a:
            ids=a["day_indices"]
            if not np.array_equal(ids,np.arange(31,365)):
                raise AssertionError(f"{name}: incomplete dates")
            q,c,d,s,e,u=(a[k] for k in ("Q","C","D","S","E","U"))
            for x in (q,c,d,e,u):
                if x.shape!=(334,144) or not np.isfinite(x).all() or x.min() < -1e-5:
                    raise AssertionError(f"{name}: invalid trajectory")
            if s.shape!=(334,145) or not np.isfinite(s).all():
                raise AssertionError("Invalid state shape")
            checks={
                "bus_balance_kwh":close("bus",q+e+dataset.pv_kw[ids]/6+d-dataset.load_kw[ids]/6-c-u),
                "battery_state_kwh":close("state",np.diff(s,axis=1)-config["eta_charge"]*c+d/config["eta_discharge"]),
                "cross_day_kwh":close("cross day",s[1:,0]-s[:-1,-1]),
                "initial_kwh":close("initial",s[0,0]-6000),
                "year_end_kwh":close("year end",s[-1,-1]-6000),
                "storage_bounds_kwh":close("bounds",max(0,1200-s.min(),s.max()-10800)),
                "power_kwh":close("power",max(0,c.max()-5000/6,d.max()-5000/6)),
                "simultaneous_kwh":close("simultaneous",np.minimum(np.maximum(c,0),np.maximum(d,0))),
                "emergency_rule_kwh":close("E",e-np.maximum(dataset.load_kw[ids]/6+c-q-dataset.pv_kw[ids]/6-d,0)),
                "unused_rule_kwh":close("U",u-np.maximum(q+dataset.pv_kw[ids]/6+d-dataset.load_kw[ids]/6-c,0)),
            }
            planned=q@dataset.prices
            emergency=e@(5*dataset.prices)
            cash=planned+emergency
            checks["annual_cash_reconciliation_yuan"]=close("annual cash",cash.sum()-summary["total_cash_cost_yuan"],1e-4)
            with (folder/"daily_results.csv").open(encoding="utf-8-sig",newline="") as stream:
                records=list(csv.DictReader(stream))
            if [r["date"] for r in records]!=[dataset.dates[i].isoformat() for i in ids]:
                raise AssertionError("Daily results dates are not aligned")
            checks["daily_cash_reconciliation_yuan"]=close("daily cash",cash-np.array([float(r["cash_cost_yuan"]) for r in records]))
            info=json.loads((folder/"information_audit.json").read_text(encoding="utf-8"))
            if len(info)!=334:
                raise AssertionError("Missing information audit records")
            for audit,day in zip(info,ids):
                origin=dataset.dates[day].isoformat()
                if audit["origin"]!=origin or audit["training_end"]>=origin or (audit["residual_end"] is not None and audit["residual_end"]>=origin):
                    raise AssertionError("Information boundary violated")
            results[name]={"days":334,"checks":checks,"cash_recomputed_yuan":float(cash.sum())}
    return results


def audit_workbooks(run,dataset):
    q1=json.loads((run/"q1.json").read_text(encoding="utf-8"))["milp"]["plan"]
    checks={}
    w=openpyxl.load_workbook(run/"deliverables/result1.xlsx",read_only=True,data_only=True)
    if w.sheetnames!=["计划购电量","充放电量"]:
        raise AssertionError("Q1 template sheets changed unexpectedly")
    rows=list(w["计划购电量"].iter_rows(min_row=2,max_row=145,max_col=2,values_only=True))
    if [r[0] for r in rows]!=[interval_label(i) for i in range(144)]:
        raise AssertionError("Q1 export time mapping incorrect")
    checks["q1_purchase_kwh"]=close("q1 Q",np.array([r[1] for r in rows])-q1["Q"])
    rows=list(w["充放电量"].iter_rows(min_row=2,max_row=7,max_col=5,values_only=True))
    checks["q1_charge_blocks_kwh"]=close("q1 C blocks",np.array([r[1] for r in rows])-np.array(q1["C"]).reshape(6,24).sum(axis=1))
    checks["q1_discharge_blocks_kwh"]=close("q1 D blocks",np.array([r[2] for r in rows])-np.array(q1["D"]).reshape(6,24).sum(axis=1))
    close("q1 start",rows[0][4]-6000);close("q1 end",rows[1][4]-6000)
    w.close()
    a=np.load(run/"main_weighted28/trajectories.npz")
    dates=[dataset.dates[i] for i in a["day_indices"]]
    w=openpyxl.load_workbook(run/"deliverables/result2.xlsx",read_only=True,data_only=True)
    if w.sheetnames!=["计划购电量","充放电量","紧急购电量"]:
        raise AssertionError("Q2 template sheets changed unexpectedly")
    p=w["计划购电量"]
    if [p.cell(1,t+2).value for t in range(144)]!=[interval_label(t) for t in range(144)]:
        raise AssertionError("Q2 time header mismatch")
    rows=list(p.iter_rows(min_row=2,max_row=335,max_col=147,values_only=True))
    if any(not isinstance(r[0],datetime) or r[0].date()!=day for r,day in zip(rows,dates)):
        raise AssertionError("Q2 dates are not typed/aligned")
    numbers=np.array([r[1:] for r in rows],dtype=float)
    checks["q2_purchase_matrix_kwh"]=close("q2 Q",numbers[:,:144]-a["Q"])
    checks["q2_daily_purchase_sums_kwh"]=close("q2 sums",numbers[:,144]-a["Q"].sum(axis=1))
    checks["q2_daily_cash_yuan"]=close("q2 cash",numbers[:,145]-(a["Q"]@dataset.prices+a["E"]@(5*dataset.prices)))
    rows=list(w["充放电量"].iter_rows(min_row=2,max_row=2005,max_col=6,values_only=True))
    if len(rows)!=334*6:
        raise AssertionError("Q2 battery row count")
    for i,day in enumerate(dates):
        block=rows[i*6:(i+1)*6]
        if not isinstance(block[0][0],datetime) or block[0][0].date()!=day:
            raise AssertionError("Battery block date mismatch")
        close("block C",np.array([r[2] for r in block])-a["C"][i].reshape(6,24).sum(axis=1))
        close("block D",np.array([r[3] for r in block])-a["D"][i].reshape(6,24).sum(axis=1))
        close("block S start",block[0][5]-a["S"][i,0]);close("block S end",block[1][5]-a["S"][i,-1])
    rows=[r for r in w["紧急购电量"].iter_rows(min_row=2,max_col=3,values_only=True) if any(x is not None for x in r)]
    totals={day:0. for day in dates};current=None;last_end=-1
    for row in rows:
        if row[0] is not None:
            if not isinstance(row[0],datetime):
                raise AssertionError("Emergency dates must be typed")
            current=row[0].date();last_end=-1
        if current not in totals or not isinstance(row[1],str) or not isinstance(row[2],(int,float)):
            raise AssertionError("Invalid emergency row")
        lo,hi=row[1].split("-")
        start=int(lo[:2])*6+int(lo[3:])//10
        end=int(hi[:2])*6+int(hi[3:])//10
        if not 0<=start<end<=144 or start<last_end:
            raise AssertionError("Emergency intervals overlap or cross day incorrectly")
        close("event sum",row[2]-a["E"][dates.index(current),start:end].sum())
        totals[current]+=row[2];last_end=end
    checks["emergency_export_daily_kwh"]=close("E export",np.array([totals[day] for day in dates])-a["E"].sum(axis=1),1e-4)
    for sheet in w:
        if any(cell.data_type=="e" for row in sheet for cell in row):
            raise AssertionError("Formula error in exported workbook")
    w.close();a.close()
    checks["exported_dates"]=334;checks["emergency_segments"]=len(rows)
    return checks


def main():
    parser=argparse.ArgumentParser()
    parser.add_argument("run",type=Path)
    parser.add_argument("--workbooks",action="store_true")
    args=parser.parse_args()
    dataset=read_dataset(ROOT/"data/processed")
    metadata=json.loads((args.run/"config.json").read_text(encoding="utf-8"))
    if dataset.sources!=metadata["input_sha256"]:
        raise AssertionError("Processed input hashes changed")
    manifest=json.loads((ROOT/"data/raw/manifest.json").read_text(encoding="utf-8"))
    source=Path(manifest["source_directory"])
    for item in manifest["files"]:
        for path in (ROOT/"data/raw"/item["file"],source/item["file"]):
            if hashlib.sha256(path.read_bytes()).hexdigest()!=item["sha256"]:
                raise AssertionError(f"Source or snapshot changed: {path}")
    report={"raw_files_unchanged":len(manifest["files"]),"processed_hashes_match":True,
            "experiments":audit_arrays(args.run,dataset,metadata["config"])}
    if args.workbooks:
        report["workbooks"]=audit_workbooks(args.run,dataset)
    report["status"]="passed"
    (args.run/"independent_audit.json").write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding="utf-8")
    print(json.dumps({"status":"passed","experiments":len(report["experiments"]),"workbooks_checked":args.workbooks},ensure_ascii=False))


if __name__=="__main__":
    main()
